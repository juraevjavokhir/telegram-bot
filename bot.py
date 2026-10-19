"""
Telegram-бот "Заявка для Китая"
================================
Собирает у пользователя список запчастей и по кнопке "Готово" формирует
Excel-файл заявки (колонки: №, Фото, Наименование, Артикул, Количество,
Ед. изм., Цена, Сумма) и сразу отправляет его в чат.

Быстрые способы добавить позиции:
  * фото с подписью «Наименование / Артикул / Кол-во [/ Ед.]» — одним сообщением;
  * текст списком, по строке на позицию: «Наименование; Артикул; Кол-во; Ед.»
    (можно вставить строки, скопированные из Excel);
  * «Артикул Кол-во» — если артикул уже встречался, название подставится само;
  * альбом фото без подписей — бот по очереди спросит недостающее;
  * Excel прошлой заявки — все позиции (с фото) загружаются в текущую.

Запуск:
    pip install -r requirements.txt
    export BOT_TOKEN="<токен от @BotFather>"
    python bot.py

Подробности — в README.md.
"""

import json
import logging
import math
import os
import re
import shutil
import uuid
from datetime import datetime

from PIL import Image as PILImage
from PIL import ImageOps
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    PicklePersistence,
    filters,
)

# --------------------------------------------------------------------------
# Настройки
# --------------------------------------------------------------------------

BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
# Если api.telegram.org заблокирован в вашей сети — задайте PROXY_URL,
# например: PROXY_URL=socks5://127.0.0.1:1080 или http://127.0.0.1:8080
PROXY_URL = os.environ.get("PROXY_URL", "")
# Здесь хранятся: незаконченные заявки (переживают перезапуск), фото,
# справочник артикулов. В Docker это volume ./data.
DATA_DIR = os.environ.get("DATA_DIR", "data")
PHOTOS_DIR = os.path.join(DATA_DIR, "photos")
CATALOG_PATH = os.path.join(DATA_DIR, "catalog.json")
STATE_PATH = os.path.join(DATA_DIR, "state.pickle")

DEFAULT_UNIT = "шт"
MAX_XLSX_BYTES = 20 * 1024 * 1024

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)
# httpx пишет в INFO каждый URL запроса, а в нём — токен бота
logging.getLogger("httpx").setLevel(logging.WARNING)


# --------------------------------------------------------------------------
# Единицы измерения, количество
# --------------------------------------------------------------------------

_UNIT_ALIASES = {
    "шт": "шт", "ш": "шт", "штук": "шт", "штука": "шт", "штуки": "шт",
    "pcs": "шт", "pc": "шт",
    "компл": "компл", "комплект": "компл", "к-т": "компл", "кт": "компл",
    "компл-т": "компл", "set": "компл",
    "м": "м", "метр": "м", "кг": "кг", "л": "л", "пара": "пара", "пар": "пара",
}


def normalize_unit(unit: str) -> str:
    """«Шт», «ШТ», «ш.» -> «шт»; «Компл» -> «компл»; прочее — в нижний регистр."""
    u = (unit or "").strip().lower().rstrip(".")
    if not u or u == "-":
        return DEFAULT_UNIT
    return _UNIT_ALIASES.get(u, u)


def _is_unit_word(token: str) -> bool:
    return token.strip().lower().rstrip(".") in _UNIT_ALIASES


def parse_qty(text: str):
    """«10», «2,5», «4 компл», «4компл» -> (кол-во, ед. или None). Иначе (None, None)."""
    m = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*([^\d\s].*)?\s*", text or "")
    if not m:
        return None, None
    qty = float(m.group(1).replace(",", "."))
    # nan/inf сюда не пройдут (только цифры), но 0 — тоже не количество
    if not math.isfinite(qty) or qty <= 0:
        return None, None
    qty = int(qty) if qty == int(qty) else qty
    unit = m.group(2).strip() if m.group(2) else None
    if unit and not _is_unit_word(unit) and len(unit) > 8:
        return None, None  # «10 штук болтов от ...» — это не количество
    return qty, (normalize_unit(unit) if unit else None)


# --------------------------------------------------------------------------
# Справочник артикулов: артикул -> наименование и ед. изм.
# Пополняется сам из каждой добавленной позиции.
# --------------------------------------------------------------------------

def article_key(article: str) -> str:
    return re.sub(r"\s+", "", article or "").upper()


def _looks_like_article(text: str) -> bool:
    t = (text or "").strip()
    return bool(t) and " " not in t and any(c.isdigit() for c in t) and len(t) >= 4


class Catalog:
    def __init__(self, path: str):
        self.path = path
        self.data: dict[str, dict] = {}
        try:
            with open(path, encoding="utf-8") as f:
                self.data = json.load(f)
        except FileNotFoundError:
            pass
        except Exception:
            logger.exception("Не удалось прочитать справочник %s", path)

    def get(self, article: str):
        key = article_key(article)
        return self.data.get(key) if key else None

    def learn(self, item: dict) -> None:
        key = article_key(item.get("article", ""))
        if not key or not item.get("name"):
            return
        new = {"article": item["article"], "name": item["name"], "unit": item.get("unit", DEFAULT_UNIT)}
        if self.data.get(key) == new:
            return
        self.data[key] = new
        try:
            os.makedirs(os.path.dirname(self.path) or ".", exist_ok=True)
            tmp = self.path + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(self.data, f, ensure_ascii=False, indent=1)
            os.replace(tmp, self.path)
        except Exception:
            logger.exception("Не удалось сохранить справочник")


CATALOG = Catalog(CATALOG_PATH)


# --------------------------------------------------------------------------
# Разбор текста в позицию
# --------------------------------------------------------------------------

def _apply_catalog(d: dict) -> dict:
    """Если артикул известен — подставляем наименование и ед. изм."""
    known = CATALOG.get(d.get("article", ""))
    if known:
        d.setdefault("name", known["name"])
        d.setdefault("unit", known.get("unit", DEFAULT_UNIT))
        d["article"] = known["article"]
    return d


def parse_fields(fields: list[str]) -> dict:
    """Поля в порядке: Наименование, Артикул, Кол-во, Ед. изм."""
    fields = [f.strip() for f in fields]
    # Строка, скопированная из Excel вместе с колонками «№» и «Фото»
    if len(fields) >= 5 and fields[0].isdigit():
        fields = fields[1:]
    if len(fields) >= 5 and fields[0] == "":
        fields = fields[1:]
    name, article, qty_raw, unit = (fields + ["", "", "", ""])[:4]
    if article == "-":
        article = ""

    d: dict = {"article": article}
    if name:
        d["name"] = name
    qty, qty_unit = parse_qty(qty_raw)
    if qty is not None:
        d["qty"] = qty
    if unit:
        d["unit"] = normalize_unit(unit)
    elif qty_unit:
        d["unit"] = qty_unit

    # Вписали артикул в поле наименования
    if not article and name and CATALOG.get(name):
        d.pop("name")
        d["article"] = name
    return _apply_catalog(d)


def parse_line(text: str) -> dict:
    """Одна позиция из строки текста. Возвращает то, что удалось понять."""
    s = (text or "").strip()
    if not s:
        return {}
    if ";" in s or "\t" in s:
        return parse_fields(re.split(r"[;\t]", s))

    # Без разделителей: «A3463530435 2», «A3463530435 4 компл», «Полуось», «2916010-2001»
    tokens = s.split()
    d: dict = {}
    if len(tokens) >= 3 and _is_unit_word(tokens[-1]):
        qty, _ = parse_qty(tokens[-2])
        if qty is not None:
            d["qty"], d["unit"] = qty, normalize_unit(tokens[-1])
            tokens = tokens[:-2]
    if "qty" not in d and len(tokens) >= 2:
        qty, unit = parse_qty(tokens[-1])
        if qty is not None:
            d["qty"] = qty
            if unit:
                d["unit"] = unit
            tokens = tokens[:-1]
    rest = " ".join(tokens)

    if CATALOG.get(rest):
        d["article"] = rest
    elif "qty" in d and _looks_like_article(rest):
        d["article"] = rest  # новый артикул — спросим наименование
    else:
        d["name"] = rest
    return _apply_catalog(d)


def parse_caption(caption: str) -> dict:
    """Подпись к фото: построчно (наименование / артикул / кол-во / ед.) или одной строкой."""
    lines = [ln.strip() for ln in (caption or "").splitlines() if ln.strip()]
    if not lines:
        return {}
    if len(lines) == 1:
        return parse_line(lines[0])
    return parse_fields(lines)


def missing_field(d: dict):
    for field in ("name", "article", "qty"):
        if field not in d:
            return field
    return None


# --------------------------------------------------------------------------
# Сессия пользователя (хранится в context.user_data, переживает перезапуск)
#   items  — готовые позиции
#   draft  — позиция, по которой сейчас задаём вопросы (или None)
#   queue  — ещё не разобранные фото (например, из альбома)
# --------------------------------------------------------------------------

def _session(context: ContextTypes.DEFAULT_TYPE) -> dict:
    ud = context.user_data
    ud.setdefault("items", [])
    ud.setdefault("queue", [])
    ud.setdefault("draft", None)
    if not ud.get("tmp_dir") or not os.path.isdir(ud["tmp_dir"]):
        ud["tmp_dir"] = os.path.join(PHOTOS_DIR, uuid.uuid4().hex)
        os.makedirs(ud["tmp_dir"], exist_ok=True)
    return ud


def _remove_photo(item) -> None:
    path = (item or {}).get("photo_path")
    if path:
        for p in (path, path + ".thumb.jpg"):
            try:
                os.remove(p)
            except OSError:
                pass


def _cleanup(context: ContextTypes.DEFAULT_TYPE) -> None:
    """Удаляет временные файлы (фото) и очищает сессию."""
    tmp_dir = context.user_data.get("tmp_dir")
    if tmp_dir and os.path.isdir(tmp_dir):
        shutil.rmtree(tmp_dir, ignore_errors=True)
    context.user_data.clear()


def add_item(session: dict, item: dict) -> str:
    """Добавляет готовую позицию; одинаковый артикул — суммирует количество.
    Возвращает короткое описание для ответа пользователю."""
    item = dict(item)
    item["unit"] = normalize_unit(item.get("unit", ""))
    item["article"] = (item.get("article") or "").strip()
    CATALOG.learn(item)

    key = article_key(item["article"])
    if key:
        for i, old in enumerate(session["items"], start=1):
            if article_key(old.get("article", "")) == key:
                total = old["qty"] + item["qty"]
                old["qty"] = int(total) if total == int(total) else round(total, 3)
                if not old.get("photo_path") and item.get("photo_path"):
                    old["photo_path"] = item["photo_path"]
                elif item.get("photo_path") and item["photo_path"] != old.get("photo_path"):
                    _remove_photo(item)  # у позиции уже есть фото — второе не нужно
                return f"⚠️ Артикул {item['article']} уже есть (№{i}) — количество теперь {old['qty']} {old['unit']}"
    session["items"].append(item)
    return f"✅ №{len(session['items'])}: {item_line(item)}"


def item_line(item: dict) -> str:
    art = f" — {item['article']}" if item.get("article") else ""
    photo = " 📷" if item.get("photo_path") else ""
    return f"{item['name']}{art} — {item['qty']} {item['unit']}{photo}"


DONE_KEYBOARD = InlineKeyboardMarkup(
    [
        [
            InlineKeyboardButton("📋 Список", callback_data="show_list"),
            InlineKeyboardButton("✅ Готово", callback_data="finish"),
        ]
    ]
)

QUESTIONS = {
    "name": "Введите наименование детали\n(или артикул, если эту деталь уже заказывали):",
    "article": "Введите артикул. Если артикула нет — отправьте «-».",
    "qty": "Сколько нужно? Например: 10  или  4 компл",
}

HELP_TEXT = (
    "Собираем заявку для партнёра из Китая. Быстрые способы:\n\n"
    "📷 <b>Фото с подписью</b> — позиция добавится сразу:\n"
    "<code>Полуось\nA3463530435\n2</code>\n\n"
    "📝 <b>Списком без фото</b> — по строке на позицию (можно вставить из Excel):\n"
    "<code>Полуось; A3463530435; 2\nЗамок с тросом; 5302500B45C00; 2; компл</code>\n\n"
    "🔁 <b>Уже заказывали</b> — хватит артикула и количества:\n"
    "<code>A3463530435 2</code>\n\n"
    "🖼 <b>Несколько фото альбомом</b> без подписей — спрошу по каждому.\n"
    "📎 <b>Excel прошлой заявки</b> — отправьте файлом, позиции загрузятся с фото.\n\n"
    "Единица по умолчанию — «шт». Одинаковые артикулы складываются.\n\n"
    "/list — список, /qty 5 10 — у №5 поставить 10, /del 5 — удалить №5,\n/skip — пропустить текущее фото, "
    "/done — сформировать Excel, /cancel — начать заново."
)


# --------------------------------------------------------------------------
# Формирование и отправка Excel
# --------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1FA398", end_color="1FA398", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

HEADERS = ["№", "Фото", "Наименование", "Артикул", "Количество", "Ед. изм.", "Цена", "Сумма"]
COL_WIDTHS = {"A": 6, "C": 34, "D": 16, "E": 12, "F": 10, "G": 12, "H": 12}
N_COLS = len(HEADERS)  # 8

# Точный размер фото в ячейке: 3 см (ширина) x 2.5 см (высота)
CM_TO_PX = 96 / 2.54          # пикселей в 1 см при 96 dpi (Excel/openpyxl)
CM_TO_PT = 72 / 2.54          # points в 1 см (высота строки задаётся в points)
PHOTO_W_CM = 3.0
PHOTO_H_CM = 2.5
PHOTO_W_PX = PHOTO_W_CM * CM_TO_PX          # ~113.4 px — реальный размер в Excel
PHOTO_H_PX = PHOTO_H_CM * CM_TO_PX          # ~94.5 px
ROW_HEIGHT_PT = PHOTO_H_CM * CM_TO_PT       # ~70.9 pt — высота строки под фото
COL_B_WIDTH = (PHOTO_W_PX - 5) / 7          # ширина колонки B в Excel-единицах
UPSCALE = 4                                  # запас разрешения для чёткости при печати/зуме
THUMB_W_PX = round(PHOTO_W_PX * UPSCALE)
THUMB_H_PX = round(PHOTO_H_PX * UPSCALE)


def build_workbook(items: list[dict], out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявка"

    # Заголовок листа
    last_col_letter = get_column_letter(N_COLS)
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = f"Заявка на запчасти — {datetime.now().strftime('%d.%m.%Y')}"
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 24

    header_row = 2
    for col_idx, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width
    ws.column_dimensions["B"].width = COL_B_WIDTH  # под фото ровно 3 см

    first_data_row = header_row + 1
    for i, item in enumerate(items):
        row = first_data_row + i

        ws.cell(row=row, column=1, value=i + 1)
        # B: фото — заполняется картинкой ниже
        ws.cell(row=row, column=3, value=item["name"])
        ws.cell(row=row, column=4, value=item.get("article", ""))
        ws.cell(row=row, column=5, value=item["qty"])
        ws.cell(row=row, column=6, value=item["unit"])
        # G: цена — пусто, заполнит партнёр
        ws.cell(row=row, column=7, value=None)
        # H: сумма — формула, посчитается сама, когда впишут цену
        ws.cell(row=row, column=8, value=f"=E{row}*G{row}")

        for col in range(1, N_COLS + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = BORDER
            cell.alignment = CENTER if col != 3 else LEFT
            cell.font = BODY_FONT
        ws.cell(row=row, column=7).number_format = "#,##0.00"
        ws.cell(row=row, column=8).number_format = "#,##0.00"

        ws.row_dimensions[row].height = ROW_HEIGHT_PT

        # Вставляем фото ровно 3x2.5 см, кадрируя по центру под нужные
        # пропорции и сохраняя высокое разрешение для чёткости при печати.
        photo_path = item.get("photo_path")
        if photo_path and os.path.exists(photo_path):
            try:
                with PILImage.open(photo_path) as im:
                    im = ImageOps.exif_transpose(im).convert("RGB")
                    im = ImageOps.fit(
                        im,
                        (THUMB_W_PX, THUMB_H_PX),
                        method=PILImage.LANCZOS,
                    )
                    resized_path = photo_path + ".thumb.jpg"
                    im.save(resized_path, format="JPEG", quality=95)
                xl_img = XLImage(resized_path)
                xl_img.width = PHOTO_W_PX
                xl_img.height = PHOTO_H_PX
                ws.add_image(xl_img, f"B{row}")
            except Exception:
                logger.exception("Не удалось вставить фото для строки %s", row)

    # Итоговая строка
    total_row = first_data_row + len(items)
    ws.cell(row=total_row, column=3, value="Итого:")
    ws.cell(row=total_row, column=3).font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=total_row, column=3).alignment = LEFT
    total_cell = ws.cell(
        row=total_row,
        column=8,
        value=f"=SUM(H{first_data_row}:H{total_row - 1})" if items else 0,
    )
    total_cell.font = Font(name="Arial", bold=True, size=10)
    total_cell.number_format = "#,##0.00"
    for col in range(1, N_COLS + 1):
        ws.cell(row=total_row, column=col).border = BORDER

    # Легенда для партнёра
    legend_row = total_row + 2
    ws.cell(
        row=legend_row,
        column=1,
        value="Жёлтым — заполняет партнёр (цена за ед.). Сумма считается автоматически.",
    ).font = Font(name="Arial", italic=True, size=9, color="808080")
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=N_COLS)

    yellow = PatternFill(start_color="FFF2A8", end_color="FFF2A8", fill_type="solid")
    for i in range(len(items)):
        ws.cell(row=first_data_row + i, column=7).fill = yellow

    ws.freeze_panes = "A3"
    wb.save(out_path)


# --------------------------------------------------------------------------
# Загрузка позиций из Excel (прошлая заявка или любой файл с такими колонками)
# --------------------------------------------------------------------------

def import_xlsx(path: str, photo_dir: str) -> tuple[list[dict], list[str]]:
    """Читает позиции из xlsx. Ищет строку-шапку с «Наименование»;
    фото берёт из картинок, привязанных к строкам. Возвращает (позиции, ошибки)."""
    import io

    wb = load_workbook(path)
    ws = wb.active

    header_row, cols = None, {}
    for r in range(1, min(ws.max_row, 15) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            v = str(ws.cell(row=r, column=c).value or "").strip().lower()
            if not v:
                continue
            if "наимен" in v:
                found.setdefault("name", c)
            elif "артикул" in v or "код" in v:
                found.setdefault("article", c)
            elif v.startswith("кол"):
                found.setdefault("qty", c)
            elif v.startswith("ед"):
                found.setdefault("unit", c)
            elif "фото" in v:
                found.setdefault("photo", c)
        if "name" in found and "qty" in found:
            header_row, cols = r, found
            break
    if header_row is None:
        raise ValueError("не нашёл шапку с колонками «Наименование» и «Количество»")

    images = {}
    for img in getattr(ws, "_images", []):
        fr = getattr(img.anchor, "_from", None)
        if fr is None:
            continue
        if "photo" in cols and fr.col + 1 != cols["photo"]:
            continue
        images.setdefault(fr.row + 1, img)

    def cell(r, key):
        c = cols.get(key)
        v = ws.cell(row=r, column=c).value if c else None
        if isinstance(v, float) and v == int(v):
            v = int(v)
        return "" if v is None else str(v).strip()

    items, errors = [], []
    for r in range(header_row + 1, ws.max_row + 1):
        name, article = cell(r, "name"), cell(r, "article")
        if name.lower().startswith("итого"):
            break
        if not name and not article:
            continue
        qty, qty_unit = parse_qty(cell(r, "qty"))
        if qty is None:
            errors.append(f"строка {r}: «{name or article}» — нет количества")
            continue
        item = {"name": name, "article": article, "qty": qty,
                "unit": cell(r, "unit") or qty_unit or DEFAULT_UNIT}
        if not name:
            known = CATALOG.get(article)
            if not known:
                errors.append(f"строка {r}: артикул {article} — нет наименования")
                continue
            item["name"] = known["name"]
        img = images.get(r)
        if img is not None:
            try:
                with PILImage.open(io.BytesIO(img._data())) as im:
                    p = os.path.join(photo_dir, f"{uuid.uuid4().hex}.jpg")
                    im.convert("RGB").save(p, format="JPEG", quality=92)
                item["photo_path"] = p
            except Exception:
                logger.exception("Не удалось извлечь фото из строки %s", r)
        items.append(item)
    return items, errors


# --------------------------------------------------------------------------
# Диалог
# --------------------------------------------------------------------------

async def _send_long(message, lines: list[str], **kwargs) -> None:
    """Отправляет список строк, разбивая на сообщения до 4000 символов."""
    chunk = ""
    chunks = []
    for line in lines:
        if len(chunk) + len(line) + 1 > 4000:
            chunks.append(chunk)
            chunk = ""
        chunk += line + "\n"
    if chunk:
        chunks.append(chunk)
    for i, text in enumerate(chunks):
        extra = kwargs if i == len(chunks) - 1 else {}
        await message.reply_text(text, **extra)


async def ask_next(message, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Задаёт вопрос по текущей позиции; если всё известно — добавляет её
    и переходит к следующему фото из очереди."""
    session = _session(context)
    while True:
        draft = session["draft"]
        if draft is None:
            if not session["queue"]:
                return
            draft = session["draft"] = session["queue"].pop(0)

        field = missing_field(draft)
        if field is None:
            item = {k: v for k, v in draft.items() if not k.startswith("_") and k != "file_id"}
            text = add_item(session, item)
            session["draft"] = None
            if session["queue"]:
                await message.reply_text(text)
            else:
                await message.reply_text(
                    f"{text}\nВсего позиций: {len(session['items'])}. Присылайте следующую.",
                    reply_markup=DONE_KEYBOARD,
                )
            continue

        question = QUESTIONS[field]
        if draft.get("_show_photo") and draft.get("file_id"):
            draft["_show_photo"] = False
            left = len(session["queue"])
            note = f"Это фото (в очереди ещё {left}).\n\n" if left else "Это фото.\n\n"
            await message.reply_photo(draft["file_id"], caption=note + question)
        else:
            await message.reply_text(question)
        return


def answer_field(draft: dict, field: str, text: str):
    """Записывает ответ на заданный вопрос. Возвращает текст ошибки или None."""
    t = (text or "").strip()
    if not t:
        return "Пустой ответ. " + QUESTIONS[field]
    if ";" in t or "\t" in t:  # ответили сразу несколькими полями
        for k, v in parse_line(t).items():
            draft.setdefault(k, v)
        return None
    if field == "name":
        if CATALOG.get(t) and "article" not in draft:
            draft["article"] = t
            _apply_catalog(draft)
        else:
            draft["name"] = t
    elif field == "article":
        draft["article"] = "" if t == "-" else t
        _apply_catalog(draft)
    elif field == "qty":
        qty, unit = parse_qty(t)
        if qty is None:
            return "Не похоже на количество. Введите число, например: 10"
        draft["qty"] = qty
        if unit:
            draft["unit"] = unit
    return None


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    session = _session(context)
    await update.message.reply_html(HELP_TEXT)
    if session["items"] or session["draft"] or session["queue"]:
        await update.message.reply_text(
            f"У вас незаконченная заявка: позиций {len(session['items'])}. "
            "Продолжайте добавлять, /done — сформировать Excel, /cancel — удалить и начать заново."
        )


async def receive_photo(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    session = _session(context)
    msg = update.message
    photo = msg.photo[-1]  # самое большое разрешение
    tg_file = await photo.get_file()
    local_path = os.path.join(session["tmp_dir"], f"{uuid.uuid4().hex}.jpg")
    await tg_file.download_to_drive(local_path)

    d = parse_caption(msg.caption or "")
    d["photo_path"] = local_path
    d["file_id"] = photo.file_id
    in_album = bool(msg.media_group_id)
    # В альбоме показываем фото вместе с вопросом, чтобы было понятно, о каком речь
    d["_show_photo"] = in_album

    if session["draft"] is not None:
        d["_show_photo"] = True
        session["queue"].append(d)
        if not in_album:
            await msg.reply_text(
                f"Фото поставлено в очередь ({len(session['queue'])}). "
                "Сначала ответьте на вопрос выше (или /skip)."
            )
        return
    session["draft"] = d
    await ask_next(msg, context)


async def receive_text(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    session = _session(context)
    msg = update.message
    text = msg.text or ""

    draft = session["draft"]
    if draft is not None:
        error = answer_field(draft, missing_field(draft), text)
        if error:
            await msg.reply_text(error)
            return
        await ask_next(msg, context)
        return

    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) > 1:
        report, errors = [], []
        for n, line in enumerate(lines, start=1):
            d = parse_line(line)
            if "name" in d and "qty" in d:
                d.setdefault("article", "")
                report.append(add_item(session, d))
            else:
                lack = "наименования" if "name" not in d else "количества"
                errors.append(f"❌ строка {n}: «{line.strip()[:40]}» — не хватает {lack}")
        if errors:
            errors.append("\nИсправьте эти строки и отправьте их ещё раз.")
        report.append(f"\nВсего позиций: {len(session['items'])}.")
        await _send_long(msg, report + errors, reply_markup=DONE_KEYBOARD)
        return

    session["draft"] = parse_line(text)
    await ask_next(msg, context)


async def receive_document(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    session = _session(context)
    msg = update.message
    doc = msg.document
    if not (doc.file_name or "").lower().endswith(".xlsx"):
        await msg.reply_text(
            "Принимаю только Excel (.xlsx). Фото отправляйте как фото, не файлом."
        )
        return
    if doc.file_size and doc.file_size > MAX_XLSX_BYTES:
        await msg.reply_text("Файл слишком большой (больше 20 МБ).")
        return

    path = os.path.join(session["tmp_dir"], f"import_{uuid.uuid4().hex}.xlsx")
    await (await doc.get_file()).download_to_drive(path)
    try:
        items, errors = import_xlsx(path, session["tmp_dir"])
    except Exception as e:
        logger.exception("Ошибка импорта xlsx")
        await msg.reply_text(f"Не получилось прочитать файл: {e}")
        return
    finally:
        try:
            os.remove(path)
        except OSError:
            pass

    report = [add_item(session, it) for it in items]
    report.insert(0, f"Загружено из файла: {len(items)} поз.")
    report.append(f"\nВсего позиций: {len(session['items'])}. "
                  "Поменять количество: /qty номер кол-во, удалить: /del номер.")
    await _send_long(msg, report + [f"❌ {e}" for e in errors], reply_markup=DONE_KEYBOARD)


async def show_list(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    session = _session(context)
    message = update.effective_message
    if update.callback_query:
        await update.callback_query.answer()
    items = session["items"]
    if not items:
        await message.reply_text("Заявка пока пустая. Пришлите фото или строку с позицией.")
        return
    lines = [f"{i}. {item_line(it)}" for i, it in enumerate(items, start=1)]
    pending = len(session["queue"]) + (1 if session["draft"] else 0)
    if pending:
        lines.append(f"\n⏳ Ещё не заполнено фото: {pending}")
    lines.append("\nПоменять количество: /qty 3 10 · удалить: /del 3")
    await _send_long(message, lines, reply_markup=DONE_KEYBOARD)


async def delete_item(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    session = _session(context)
    items = session["items"]
    try:
        n = int(context.args[0])
        if not 1 <= n <= len(items):
            raise ValueError
    except (IndexError, ValueError, TypeError):
        await update.message.reply_text(
            f"Укажите номер от 1 до {len(items)}, например: /del 3" if items
            else "Заявка пустая — удалять нечего."
        )
        return
    item = items.pop(n - 1)
    _remove_photo(item)
    await update.message.reply_text(
        f"🗑 Удалено №{n}: {item_line(item)}\nОсталось позиций: {len(items)}. Номера сдвинулись — /list"
    )


async def set_qty(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """/qty 5 10 — поменять количество у позиции №5 на 10."""
    session = _session(context)
    items = session["items"]
    try:
        n = int(context.args[0])
        qty, unit = parse_qty(" ".join(context.args[1:]))
        if not 1 <= n <= len(items) or qty is None:
            raise ValueError
    except (IndexError, ValueError, TypeError):
        await update.message.reply_text("Формат: /qty номер количество, например: /qty 5 10")
        return
    item = items[n - 1]
    item["qty"] = qty
    if unit:
        item["unit"] = unit
    await update.message.reply_text(f"✏️ №{n}: {item_line(item)}")


async def skip(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    session = _session(context)
    if session["draft"] is None:
        await update.message.reply_text("Сейчас нечего пропускать.")
        return
    _remove_photo(session["draft"])
    session["draft"] = None
    await update.message.reply_text("Пропущено.")
    if session["queue"]:
        await ask_next(update.message, context)


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    _cleanup(context)
    await update.message.reply_text("Заявка удалена. Можно начинать новую — присылайте позиции.")


async def add_more(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Кнопка из старой версии бота — на случай нажатия на старые сообщения."""
    await update.callback_query.answer()
    await update.effective_message.reply_text("Присылайте следующую позицию.")


async def finish(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Кнопка «✅ Готово» или команда /done."""
    message = update.effective_message
    if update.callback_query:
        await update.callback_query.answer()
        try:
            await update.callback_query.edit_message_reply_markup(reply_markup=None)
        except Exception:
            pass  # сообщение уже без кнопок

    session = _session(context)
    items = session["items"]
    if not items:
        await message.reply_text("Список пуст — нечего сохранять. Пришлите фото или строку с позицией.")
        return

    pending = len(session["queue"]) + (1 if session["draft"] else 0)
    await message.reply_text("Формирую Excel-файл…")

    out_path = os.path.join(session["tmp_dir"], "zayavka.xlsx")
    build_workbook(items, out_path)

    caption = f"Готово ✅ Позиций: {len(items)}"
    if pending:
        caption += f"\n⚠️ Не вошли недозаполненные фото: {pending}"
    date_str = datetime.now().strftime("%Y-%m-%d")
    with open(out_path, "rb") as f:
        await message.reply_document(
            document=f,
            filename=f"Заявка_{date_str}.xlsx",
            caption=caption,
        )

    _cleanup(context)


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------

def main() -> None:
    if not BOT_TOKEN:
        raise SystemExit(
            "Не задан BOT_TOKEN. Установите переменную окружения BOT_TOKEN "
            "(токен от @BotFather) и запустите снова."
        )
    os.makedirs(PHOTOS_DIR, exist_ok=True)

    builder = (
        Application.builder()
        .token(BOT_TOKEN)
        .persistence(PicklePersistence(filepath=STATE_PATH))
    )
    if PROXY_URL:
        logger.info("Использую прокси: %s", PROXY_URL)
        builder = builder.proxy(PROXY_URL).get_updates_proxy(PROXY_URL)
    application = builder.build()

    application.add_handler(CommandHandler(["start", "help"], start))
    application.add_handler(CommandHandler("list", show_list))
    application.add_handler(CommandHandler("del", delete_item))
    application.add_handler(CommandHandler("qty", set_qty))
    application.add_handler(CommandHandler("skip", skip))
    application.add_handler(CommandHandler("done", finish))
    application.add_handler(CommandHandler("cancel", cancel))
    application.add_handler(CallbackQueryHandler(finish, pattern="^finish$"))
    application.add_handler(CallbackQueryHandler(show_list, pattern="^show_list$"))
    application.add_handler(CallbackQueryHandler(add_more, pattern="^add_more$"))
    application.add_handler(MessageHandler(filters.PHOTO, receive_photo))
    application.add_handler(MessageHandler(filters.Document.ALL, receive_document))
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, receive_text))

    logger.info("Бот запущен, ожидаю сообщения…")
    application.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
